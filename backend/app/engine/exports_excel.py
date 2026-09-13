"""Presentation-grade multi-tab Excel (.xlsx) export generator for valuation exhibits.
"""

import io
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .models import ValuationResponse, DerivedSecurity, BreakpointTier, ClaimTierAllocation


def generate_valuation_workbook(response: ValuationResponse) -> io.BytesIO:
    """Creates a comprehensive, beautifully styled multi-tab Excel workbook
    containing all exhibits, cap tables, OPM mechanics, waterfalls, and holdings.
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default sheet

    # Design palette
    navy_fill = PatternFill(start_color="17242B", end_color="17242B", fill_type="solid")
    teal_fill = PatternFill(start_color="0B6B68", end_color="0B6B68", fill_type="solid")
    subhead_fill = PatternFill(start_color="EEF3F9", end_color="EEF3F9", fill_type="solid")
    total_fill = PatternFill(start_color="EDF6F4", end_color="EDF6F4", fill_type="solid")

    white_font_bold = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    navy_font_bold = Font(name="Calibri", size=10, bold=True, color="17242B")
    title_font = Font(name="Calibri", size=14, bold=True, color="17242B")
    regular_font = Font(name="Calibri", size=10)
    bold_font = Font(name="Calibri", size=10, bold=True)
    italic_font = Font(name="Calibri", size=9, italic=True, color="687386")

    thin_border = Border(
        left=Side(style="thin", color="DCE1E7"),
        right=Side(style="thin", color="DCE1E7"),
        top=Side(style="thin", color="DCE1E7"),
        bottom=Side(style="thin", color="DCE1E7")
    )
    total_border = Border(
        top=Side(style="thin", color="0B6B68"),
        bottom=Side(style="double", color="0B6B68")
    )

    def style_header(ws, row_idx: int, cols: int, fill=navy_fill, font=white_font_bold):
        for col_idx in range(1, cols + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(horizontal="right" if col_idx > 1 else "left", vertical="center", wrap_text=True)

    def auto_fit_columns(ws):
        ws.views.sheetView[0].showGridLines = True
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # 1. Assumptions & Index Sheet
    ws_assump = wb.create_sheet(title="Key Assumptions")
    ws_assump.append(["CONTINGENT CLAIMS ANALYSIS — VALUATION REPORT"])
    ws_assump.cell(row=1, column=1).font = title_font
    ws_assump.append([f"Company: {response.company_name} | Client: {response.client_name} | Status: {response.report_status}"])
    ws_assump.cell(row=2, column=1).font = italic_font
    ws_assump.append([])

    assumptions = [
        ("Valuation Parameter", "Calibration Date", "Valuation Date"),
        ("Effective Measurement Date", response.calibration_date, response.valuation_date),
        ("Global Expected Exit Date", response.exit_date, response.exit_date),
        ("Day Count Convention", response.day_count_name, response.day_count_name),
        ("Reporting Currency", response.report_currency, response.report_currency),
        ("Display Units", response.display_units.capitalize(), response.display_units.capitalize()),
        ("Risk-Free Rate (Annual Effective)", f"{response.rf_calibration_effective:.2f}%", f"{response.rf_valuation_effective:.2f}%"),
        ("Risk-Free Rate (Continuous rc)", f"{response.rf_calibration_continuous:.4f}%", f"{response.rf_valuation_continuous:.4f}%"),
        ("Expected Volatility (Annualized)", f"{response.calibration_opm.volatility:.1f}%", f"{response.valuation_opm.volatility:.1f}%"),
        ("Term to Liquidity (Years)", f"{response.term_calibration:.2f} yrs", f"{response.term_valuation:.2f} yrs"),
        ("Concluded Equity Value", f"{response.calibration_solved_equity:,.2f}", f"{response.concluded_equity_value:,.2f}")
    ]
    for row in assumptions:
        ws_assump.append(list(row))
    style_header(ws_assump, 4, 3)
    for r_idx in range(5, 5 + len(assumptions) - 1):
        for c_idx in range(1, 4):
            ws_assump.cell(row=r_idx, column=c_idx).border = thin_border
            ws_assump.cell(row=r_idx, column=c_idx).font = regular_font
    auto_fit_columns(ws_assump)

    # 2. Calibration & Valuation Cap Tables
    for label, secs in [("Calibration Cap Table", response.calibration_derived_securities), ("Valuation Cap Table", response.valuation_derived_securities)]:
        ws = wb.create_sheet(title=label)
        ws.append([f"Capitalization Table & Instrument Terms ({label})"])
        ws.cell(row=1, column=1).font = title_font
        ws.append([])

        headers = [
            "Share Class / Instrument", "Subtype", "Shares Outstanding", "Original Issue Price",
            "Conversion Price", "Conversion Ratio", "Seniority", "Participation",
            "Accrued Div / Share", "Total Preference Claim", "Fully Diluted Shares"
        ]
        ws.append(headers)
        style_header(ws, 3, len(headers))

        r_start = 4
        for sec in secs:
            ws.append([
                sec.security,
                sec.security_subtype,
                sec.shares,
                sec.original_issue_price if sec.original_issue_price > 0 else "—",
                sec.conversion_price if sec.conversion_price is not None else "—",
                f"{sec.conversion_ratio:.2f}x" if sec.conversion_ratio > 0 else "0.00x",
                sec.seniority if sec.seniority < 900 else "—",
                sec.participation,
                sec.per_share_dividend,
                sec.total_liquidation_preference,
                sec.fully_diluted_shares
            ])
            r_curr = ws.max_row
            for c_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=r_curr, column=c_idx)
                cell.border = thin_border
                cell.font = regular_font
                if c_idx in (3, 11):
                    cell.number_format = "#,##0"
                elif c_idx in (4, 5, 9, 10):
                    cell.number_format = "$#,##0.00"

        # Total row
        ws.append([
            "TOTAL", "",
            sum(s.shares for s in secs), "", "", "", "", "", "",
            sum(s.total_liquidation_preference for s in secs),
            sum(s.fully_diluted_shares for s in secs)
        ])
        tot_row = ws.max_row
        for c_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=tot_row, column=c_idx)
            cell.fill = total_fill
            cell.font = bold_font
            cell.border = total_border
            if c_idx in (3, 11):
                cell.number_format = "#,##0"
            elif c_idx == 10:
                cell.number_format = "$#,##0.00"

        auto_fit_columns(ws)

    # 3. Breakpoints (Valuation)
    ws_bp = wb.create_sheet(title="Breakpoint Schedule")
    ws_bp.append(["Valuation Date Breakpoint Schedule"])
    ws_bp.cell(row=1, column=1).font = title_font
    ws_bp.append([])
    bp_headers = ["Tier", "Start Equity ($)", "End Equity ($)", "Width ($)", "Claimants", "Trigger Description"]
    ws_bp.append(bp_headers)
    style_header(ws_bp, 3, len(bp_headers))
    for bp in response.valuation_breakpoints:
        ws_bp.append([
            bp.tier,
            bp.start_equity,
            bp.end_equity,
            bp.width,
            bp.claimants_description,
            bp.event_description
        ])
        r_curr = ws_bp.max_row
        for c in range(1, len(bp_headers) + 1):
            cell = ws_bp.cell(row=r_curr, column=c)
            cell.border = thin_border
            cell.font = regular_font
            if c in (2, 3, 4):
                cell.number_format = "$#,##0.00"
    auto_fit_columns(ws_bp)

    # 4. OPM Allocation (Valuation)
    ws_opm = wb.create_sheet(title="Valuation OPM Allocation")
    ws_opm.append(["Valuation Date Option Pricing Method (OPM) Allocation"])
    ws_opm.cell(row=1, column=1).font = title_font
    ws_opm.append([f"Concluded Enterprise Equity Value: ${response.concluded_equity_value:,.2f}"])
    ws_opm.append([])
    opm_headers = ["Share Class", "Shares", "Concluded Value ($)", "Value Per Share ($)", "% Allocation", "Fully Diluted %"]
    ws_opm.append(opm_headers)
    style_header(ws_opm, 4, len(opm_headers))

    total_fd_val = sum(s.fully_diluted_shares for s in response.valuation_derived_securities)
    for sec in response.valuation_derived_securities:
        name = sec.security
        tot_val = response.valuation_opm.allocated_values.get(name, 0.0)
        ps_val = response.valuation_opm.per_share_values.get(name, 0.0)
        pct_val = response.valuation_opm.percent_allocations.get(name, 0.0)
        fd_pct = (sec.fully_diluted_shares / total_fd_val) if total_fd_val > 0 else 0.0
        ws_opm.append([
            name, sec.shares, tot_val, ps_val, pct_val, fd_pct
        ])
        r_curr = ws_opm.max_row
        for c in range(1, len(opm_headers) + 1):
            cell = ws_opm.cell(row=r_curr, column=c)
            cell.border = thin_border
            cell.font = regular_font
            if c == 2:
                cell.number_format = "#,##0"
            elif c in (3, 4):
                cell.number_format = "$#,##0.00"
            elif c in (5, 6):
                cell.number_format = "0.00%"

    # Total row
    ws_opm.append([
        "TOTAL",
        sum(s.shares for s in response.valuation_derived_securities),
        response.valuation_opm.total_allocated,
        "—",
        1.0,
        1.0
    ])
    tot_opm_r = ws_opm.max_row
    for c in range(1, len(opm_headers) + 1):
        cell = ws_opm.cell(row=tot_opm_r, column=c)
        cell.fill = total_fill
        cell.font = bold_font
        cell.border = total_border
        if c == 2:
            cell.number_format = "#,##0"
        elif c == 3:
            cell.number_format = "$#,##0.00"
        elif c in (5, 6):
            cell.number_format = "0.00%"
    auto_fit_columns(ws_opm)

    # 5. Client Holdings
    ws_hold = wb.create_sheet(title="Client Holdings Summary")
    ws_hold.append(["Client Holdings & Valuation Summary"])
    ws_hold.cell(row=1, column=1).font = title_font
    ws_hold.append([])
    h_headers = [
        "Fund / Vehicle", "Security", "Units Held", "Investment Cost ($)",
        "Valuation Fair Value / Share ($)", "Concluded Fair Value ($)",
        "Class Ownership %", "FD Ownership %", "MOIC"
    ]
    ws_hold.append(h_headers)
    style_header(ws_hold, 3, len(h_headers))

    for it in response.holdings.items:
        ws_hold.append([
            it.fund,
            it.security,
            it.units,
            it.cost,
            it.fair_value_per_share,
            it.concluded_fair_value,
            it.class_ownership_pct,
            it.fully_diluted_ownership_pct,
            f"{it.moic:.2f}x" if it.moic is not None else "—"
        ])
        r_curr = ws_hold.max_row
        for c in range(1, len(h_headers) + 1):
            cell = ws_hold.cell(row=r_curr, column=c)
            cell.border = thin_border
            cell.font = regular_font
            if c == 3:
                cell.number_format = "#,##0"
            elif c in (4, 5, 6):
                cell.number_format = "$#,##0.00"
            elif c in (7, 8):
                cell.number_format = "0.00%"

    # Total row
    ws_hold.append([
        "TOTAL PORTFOLIO", "",
        sum(it.units for it in response.holdings.items),
        response.holdings.total_cost,
        "—",
        response.holdings.total_value,
        "—", "—",
        f"{response.holdings.consolidated_moic:.2f}x" if response.holdings.consolidated_moic else "—"
    ])
    tot_h_r = ws_hold.max_row
    for c in range(1, len(h_headers) + 1):
        cell = ws_hold.cell(row=tot_h_r, column=c)
        cell.fill = total_fill
        cell.font = bold_font
        cell.border = total_border
        if c == 3:
            cell.number_format = "#,##0"
        elif c in (4, 6):
            cell.number_format = "$#,##0.00"

    auto_fit_columns(ws_hold)

    # 6. Waterfall Analysis
    ws_wf = wb.create_sheet(title="Waterfall Analysis")
    ws_wf.append([f"Waterfall Distribution at Selected Equity: ${response.waterfall.applied_equity:,.2f}"])
    ws_wf.cell(row=1, column=1).font = title_font
    ws_wf.append([])
    wf_headers = ["Share Class", "Shares", "Proceeds ($)", "Proceeds Per Share ($)", "% Recovery", "% of Total Proceeds"]
    ws_wf.append(wf_headers)
    style_header(ws_wf, 3, len(wf_headers))

    for item in response.waterfall.distribution:
        ws_wf.append([
            item.security,
            item.shares,
            item.proceeds,
            item.proceeds_per_share,
            item.percent_recovery,
            item.percent_of_total
        ])
        r_curr = ws_wf.max_row
        for c in range(1, len(wf_headers) + 1):
            cell = ws_wf.cell(row=r_curr, column=c)
            cell.border = thin_border
            cell.font = regular_font
            if c == 2:
                cell.number_format = "#,##0"
            elif c in (3, 4):
                cell.number_format = "$#,##0.00"
            elif c in (5, 6):
                cell.number_format = "0.00%"

    auto_fit_columns(ws_wf)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out

