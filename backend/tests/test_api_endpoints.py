"""Integration test suite for FastAPI REST API endpoints using httpx AsyncClient.
"""

import pytest
import io
import pypdf
import openpyxl
from httpx import AsyncClient, ASGITransport
from app.main import app
from app.engine.default_scenario import get_default_valuation_request


@pytest.mark.asyncio
async def test_api_health():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        resp = await client.get("/api/health")
        assert resp.status_code == 200
        data = resp.json()
        assert data["status"] == "ok"


@pytest.mark.asyncio
async def test_api_root_serves_frontend():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        resp = await client.get("/")
        assert resp.status_code == 200
        assert "<app-root>" in resp.text


@pytest.mark.asyncio
async def test_api_default_scenario_and_calculate():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        # Get default scenario
        resp_def = await client.get("/api/scenario/default")
        assert resp_def.status_code == 200
        default_req = resp_def.json()
        assert default_req["company_name"] == "TADO"

        # Validate
        resp_val = await client.post("/api/validate", json=default_req)
        assert resp_val.status_code == 200
        assert resp_val.json()["valid"] is True

        # Calculate
        resp_calc = await client.post("/api/calculate", json=default_req)
        assert resp_calc.status_code == 200
        data = resp_calc.json()
        assert data["company_name"] == "TADO"
        assert data["calibration_solved_equity"] == pytest.approx(287252502.92, abs=1.0)
        assert data["calibration_opm"]["per_share_values"]["Series I"] == pytest.approx(2021.90, abs=0.01)
        assert data["valuation_opm"]["per_share_values"]["Series I"] == pytest.approx(1992.50, abs=0.05)
        assert "calibration_waterfall" in data
        assert "comparative_waterfall" in data
        assert len(data["comparative_waterfall"]) > 0


@pytest.mark.asyncio
async def test_api_calculate_with_deleted_calibration_security():
    """Verify POST /api/calculate auto-resolves when calibration_security_name is missing or deleted."""
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        resp_def = await client.get("/api/scenario/default")
        req_data = resp_def.json()

        # Delete Series I from calibration cap table
        del req_data["calibration_securities"][0]
        # Keep calibration_security_name as 'Series I'
        req_data["calibration_security_name"] = "Series I"

        resp_calc = await client.post("/api/calculate", json=req_data)
        assert resp_calc.status_code == 200
        data = resp_calc.json()
        assert data["calibration_security_name"] == "Series H"
        assert data["calibration_solved_equity"] > 0
        assert "Series H" in data["calibration_opm"]["per_share_values"]



@pytest.mark.asyncio
async def test_api_risk_free_curves():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        resp = await client.get("/api/risk-free-rates/curves")
        assert resp.status_code == 200
        curves = resp.json()
        assert "us_treasury" in curves
        assert "ecb" in curves


@pytest.mark.asyncio
async def test_api_capital_iq_template_generation():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        payload = {
            "tickers": ["IQ247543", "IQ28472"],
            "calibration_date": "2025-02-26",
            "valuation_date": "2026-06-30",
            "currency": "EUR"
        }
        resp = await client.post("/api/capital-iq/workbook", json=payload)
        assert resp.status_code == 200
        assert len(resp.content) > 1000
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        assert "Instructions" in wb.sheetnames
        assert "Snapshot_Calibration" in wb.sheetnames


@pytest.mark.asyncio
async def test_api_export_excel():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        req = get_default_valuation_request().model_dump()
        resp = await client.post("/api/exports/excel", json=req)
        assert resp.status_code == 200
        assert len(resp.content) > 2000
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        assert "Valuation OPM Allocation" in wb.sheetnames


@pytest.mark.asyncio
async def test_api_export_pdf_with_watermark():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        req = get_default_valuation_request().model_dump()
        req["show_secondary_currency"] = True
        req["secondary_currency"] = "USD"
        req["secondary_fx_rate"] = 1.08
        resp = await client.post("/api/exports/pdf", json=req)
        assert resp.status_code == 200
        pdf_bytes = resp.content
        assert len(pdf_bytes) > 5000

        reader = pypdf.PdfReader(io.BytesIO(pdf_bytes))
        assert len(reader.pages) >= 10
        text = "".join(p.extract_text() or "" for p in reader.pages)
        assert "HIGHLY CONFIDENTIAL" in text
        assert "TADO" in text
        assert "Exhibit 1.0" in text
        assert "Exhibit 10.0" in text


